<<<<<<< Updated upstream
import os, re, json, uuid, logging, hashlib, jwt
=======
import os, re, uuid, logging, hashlib, jwt, threading
from io import BytesIO
from zipfile import BadZipFile
>>>>>>> Stashed changes
from typing import Optional
from dotenv import load_dotenv
load_dotenv()
from datetime import datetime, timedelta
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Form, BackgroundTasks, Header
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel
from sqlalchemy import create_engine, Column, Integer, String, Text, DateTime, ForeignKey, JSON, LargeBinary
from sqlalchemy import inspect as sa_inspect, text as sa_text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker, declarative_base, Session
from . import rag

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("campusflow")
SECRET = os.environ["JWT_SECRET"]
STORE = os.getenv("STORAGE_DIR", "./data/files"); os.makedirs(STORE, exist_ok=True)
PUBLIC = os.getenv("PUBLIC_URL", "http://localhost:5173")
EXT, MAX_BYTES = {"pdf", "jpg", "jpeg", "png", "docx", "txt"}, 25 << 20
<<<<<<< Updated upstream
NOTE_MAX = 100  # short note shown to students under each document
# Questions staff answer at upload time. key -> (label, placeholder example for the note)
CATEGORIES = {
    "leave":       ("Leave",        "e.g. Apply for on-duty leave by Friday 5 PM"),
    "fees":        ("Fees",         "e.g. Last date to pay semester fee is 30 Nov"),
    "competition": ("Competition",  "e.g. Inter-college coding contest, register by 5 Nov"),
    "symposium":   ("Symposium",    "e.g. National AI symposium on 12 Dec, free entry"),
    "exam":        ("Exam",         "e.g. End-semester exams start on 2 Dec"),
    "scholarship": ("Scholarship",  "e.g. Merit scholarship forms open until 20 Nov"),
    "holiday":     ("Holiday",      "e.g. College closed on 1 Nov for Kannada Rajyotsava"),
    "placement":   ("Placement",    "e.g. TCS campus drive on 8 Nov, final years only"),
    "event":       ("Event",        "e.g. Annual day on 15 Dec, all students invited"),
    "other":       ("Other",        "e.g. Short summary of what this document is about"),
}
=======
MAX_STUDENT_IMPORT_BYTES, MAX_STUDENT_IMPORT_ROWS = 5 << 20, 1000
>>>>>>> Stashed changes
NOT_FOUND = "I couldn't find this in the available published documents. Please check with the college office."

DB_URL = os.getenv("DATABASE_URL", "sqlite:///./data/campusflow.db")  # PostgreSQL URL also works
eng = create_engine(DB_URL, connect_args={"check_same_thread": False, "timeout": 30} if DB_URL.startswith("sqlite") else {}); Session_ = sessionmaker(eng); Base = declarative_base()

class User(Base):
    __tablename__ = "users"
    id = Column(Integer, primary_key=True); name = Column(String); username = Column(String, unique=True)
    email = Column(String, unique=True); pw_hash = Column(String); role = Column(String); department = Column(String)
    created_at = Column(DateTime, default=datetime.utcnow)

class Document(Base):
    __tablename__ = "documents"
    id = Column(Integer, primary_key=True); file_name = Column(String); file_type = Column(String)
    storage_path = Column(String); uploaded_by = Column(ForeignKey("users.id"))
    uploaded_by_name = Column(String); uploaded_by_role = Column(String); department = Column(String)
    uploaded_at = Column(DateTime, default=datetime.utcnow); status = Column(String, default="uploaded")
    error = Column(Text); page_count = Column(Integer)
    category = Column(String); note = Column(String(NOTE_MAX))  # staff answers given at upload

class Chunk(Base):
    __tablename__ = "document_chunks"
    id = Column(Integer, primary_key=True); document_id = Column(ForeignKey("documents.id"), index=True)
    page_number = Column(Integer); text = Column(Text); embedding = Column(LargeBinary)

class ChatSession(Base):
    __tablename__ = "chat_sessions"
    id = Column(Integer, primary_key=True); student_id = Column(ForeignKey("users.id"), index=True)
    title = Column(String); created_at = Column(DateTime, default=datetime.utcnow)
    updated_at = Column(DateTime, default=datetime.utcnow)

class Message(Base):
    __tablename__ = "chat_messages"  # meta holds sources (answer_sources), confidence, share_text
    id = Column(Integer, primary_key=True); student_id = Column(ForeignKey("users.id"), index=True)
    session_id = Column(ForeignKey("chat_sessions.id"), index=True)
    role = Column(String); text = Column(Text); meta = Column(JSON, default=dict)
    created_at = Column(DateTime, default=datetime.utcnow)

hp = lambda pw: hashlib.pbkdf2_hmac("sha256", pw.encode(), SECRET.encode(), 100_000).hex()

app = FastAPI(title="CampusFlow")
app.add_middleware(CORSMiddleware, allow_origins=[PUBLIC], allow_methods=["*"], allow_headers=["*"])

@app.on_event("startup")
def boot():
    Base.metadata.create_all(eng)
    # upgrade older databases: add session_id and group old questions into one chat
    cols = [c["name"] for c in sa_inspect(eng).get_columns("chat_messages")]
    if "session_id" not in cols:
        with eng.begin() as c: c.execute(sa_text("ALTER TABLE chat_messages ADD COLUMN session_id INTEGER"))
<<<<<<< Updated upstream
    doc_cols = [c["name"] for c in sa_inspect(eng).get_columns("documents")]
    for col, ddl in (("category", "VARCHAR"), ("note", f"VARCHAR({NOTE_MAX})")):
        if col not in doc_cols:
            with eng.begin() as c: c.execute(sa_text(f"ALTER TABLE documents ADD COLUMN {col} {ddl}"))
=======
    # upgrade: add email column if missing
    user_cols = [c["name"] for c in sa_inspect(eng).get_columns("users")]
    if "email" not in user_cols:
        with eng.begin() as c: c.execute(sa_text("ALTER TABLE users ADD COLUMN email TEXT"))
>>>>>>> Stashed changes
    with Session_() as s:
        for (sid,) in s.query(Message.student_id).filter(Message.session_id.is_(None)).distinct().all():
            cs = ChatSession(student_id=sid, title="Earlier questions"); s.add(cs); s.flush()
            s.query(Message).filter(Message.student_id == sid, Message.session_id.is_(None)).update({"session_id": cs.id})
        s.commit()
    threading.Thread(target=warm_embeddings, name="embedding-warmup", daemon=True).start()
    threading.Thread(target=warm_ollama, name="ollama-warmup", daemon=True).start()

def db():
    s = Session_()
    try: yield s
    finally: s.close()

def warm_embeddings():
    try:
        rag.emb(["CampusFlow startup warm-up"])
        log.info("embedding model warm-up complete")
    except Exception:
        log.exception("embedding model warm-up failed")

def warm_ollama():
    try:
        rag.warm_llm()
        log.info("Ollama model warm-up complete")
    except Exception:
        log.exception("Ollama model warm-up failed")

def current(authorization: str = Header(None), s: Session = Depends(db)):
    try: uid = int(jwt.decode((authorization or "").removeprefix("Bearer "), SECRET, algorithms=["HS256"])["sub"])
    except Exception: raise HTTPException(401, "Not authenticated")
    u = s.get(User, uid)
    if not u: raise HTTPException(401, "Not authenticated")
    return u

def staff(u: User = Depends(current)):
    if u.role != "staff": raise HTTPException(403, "Staff access only")
    return u

# ── Pydantic models ──────────────────────────────────────────────────────────
class Login(BaseModel): email: str; password: str; role: str
class StaffRegister(BaseModel): name: str; department: str; email: str; password: str
class StudentCreate(BaseModel): name: str; department: str; email: str; password: str
class Ask(BaseModel): question: str; session_id: Optional[int] = None

# ── Auth endpoints ───────────────────────────────────────────────────────────
@app.post("/api/register")
def register_staff(b: StaffRegister, s: Session = Depends(db)):
    """Staff self-registration."""
    if not b.name.strip() or not b.email.strip() or not b.password.strip() or not b.department.strip():
        raise HTTPException(400, "All fields are required")
    if len(b.password) < 6:
        raise HTTPException(400, "Password must be at least 6 characters")
    if s.query(User).filter_by(email=b.email.lower().strip()).first():
        raise HTTPException(409, "An account with this email already exists")
    username = b.email.lower().strip().split("@")[0]
    # ensure unique username
    base = username
    n = 1
    while s.query(User).filter_by(username=username).first():
        username = f"{base}{n}"; n += 1
    u = User(name=b.name.strip(), username=username, email=b.email.lower().strip(),
             pw_hash=hp(b.password), role="staff", department=b.department.strip())
    s.add(u); s.commit()
    log.info("staff registered: %s (%s)", u.email, u.department)
    token = jwt.encode({"sub": str(u.id), "exp": datetime.utcnow() + timedelta(hours=8)}, SECRET)
    return {"token": token}

@app.post("/api/login")
def login(b: Login, s: Session = Depends(db)):
    u = s.query(User).filter_by(email=b.email.lower().strip()).first()
    if not u or u.pw_hash != hp(b.password) or u.role != b.role:
        log.warning("failed login: %s", b.email); raise HTTPException(401, "Invalid email, password or role")
    log.info("login: %s (%s)", u.email, u.role)
    return {"token": jwt.encode({"sub": str(u.id), "exp": datetime.utcnow() + timedelta(hours=8)}, SECRET)}

@app.get("/api/me")
def me(u: User = Depends(current), s: Session = Depends(db)):
    return {"name": u.name, "role": u.role, "department": u.department, "email": u.email or "",
            "indexed": s.query(Document).filter_by(status="ready").count()}

# ── Student management (staff only) ─────────────────────────────────────────
@app.post("/api/students")
def create_student(b: StudentCreate, u: User = Depends(staff), s: Session = Depends(db)):
    """Staff creates a student account."""
    if not b.name.strip() or not b.email.strip() or not b.password.strip() or not b.department.strip():
        raise HTTPException(400, "All fields are required")
    if len(b.password) < 6:
        raise HTTPException(400, "Password must be at least 6 characters")
    if s.query(User).filter_by(email=b.email.lower().strip()).first():
        raise HTTPException(409, "An account with this email already exists")
    username = b.email.lower().strip().split("@")[0]
    base = username
    n = 1
    while s.query(User).filter_by(username=username).first():
        username = f"{base}{n}"; n += 1
    st = User(name=b.name.strip(), username=username, email=b.email.lower().strip(),
              pw_hash=hp(b.password), role="student", department=b.department.strip())
    s.add(st); s.commit()
    log.info("student created by %s: %s (%s)", u.email, st.email, st.department)
    return {"id": st.id, "name": st.name, "email": st.email, "department": st.department,
            "created_at": st.created_at.isoformat() + "Z"}

@app.post("/api/students/import")
async def import_students(file: UploadFile = File(...), u: User = Depends(staff), s: Session = Depends(db)):
    """Create student accounts from an Excel workbook; invalid rows are reported and skipped."""
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "Upload an Excel .xlsx file")
    content = await file.read(MAX_STUDENT_IMPORT_BYTES + 1)
    if len(content) > MAX_STUDENT_IMPORT_BYTES:
        raise HTTPException(413, "Excel file must be 5 MB or smaller")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as e:
        raise HTTPException(400, "Could not read the Excel file. Upload a valid .xlsx workbook.") from e

    try:
        sheet = workbook.active
        if sheet.max_row is None or sheet.max_row < 2:
            raise HTTPException(400, "The Excel sheet must include a header row and at least one student row")
        if sheet.max_row - 1 > MAX_STUDENT_IMPORT_ROWS:
            raise HTTPException(400, f"Excel file can contain at most {MAX_STUDENT_IMPORT_ROWS} student rows")

        headers = [
            re.sub(r"\s+", " ", str(cell.value or "").strip().lower())
            for cell in next(sheet.iter_rows(min_row=1, max_row=1))
        ]
        columns = {}
        for key, aliases in {
            "name": {"name"},
            "department": {"department"},
            "email": {"mail id", "email", "mail"},
            "password": {"password"},
        }.items():
            matches = [i for i, header in enumerate(headers) if header in aliases]
            if not matches:
                raise HTTPException(400, f"Excel file is missing a '{key}' column")
            if len(matches) > 1:
                raise HTTPException(400, f"Excel file has multiple columns for '{key}'")
            columns[key] = matches[0]

        def cell_text(value):
            if value is None:
                return ""
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value).strip()

        existing_emails = {
            email.lower() for (email,) in s.query(User.email).filter(User.email.isnot(None)).all()
        }
        used_usernames = {username for (username,) in s.query(User.username).all() if username}
        imported_emails = set()
        errors, created_count = [], 0

        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            fields = {key: cell_text(values[index] if index < len(values) else None)
                      for key, index in columns.items()}
            if not any(fields.values()):
                continue

            issue = None
            if any(not fields[key] for key in columns):
                issue = "Name, department, mail id, and password are required"
            elif not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", fields["email"]):
                issue = "Enter a valid email address"
            elif len(fields["password"]) < 6:
                issue = "Password must be at least 6 characters"

            email = fields["email"].lower()
            if issue is None and (email in existing_emails or email in imported_emails):
                issue = "An account with this email already exists"

            if issue:
                errors.append({"row": row_number, "message": issue})
                continue

            username = email.split("@", 1)[0]
            base_username = username
            suffix = 1
            while username in used_usernames:
                username = f"{base_username}{suffix}"
                suffix += 1
            used_usernames.add(username)
            imported_emails.add(email)
            s.add(User(
                name=fields["name"],
                username=username,
                email=email,
                pw_hash=hp(fields["password"]),
                role="student",
                department=fields["department"],
            ))
            created_count += 1

        try:
            s.commit()
        except IntegrityError as e:
            s.rollback()
            raise HTTPException(409, "An account with one of these email addresses already exists. No accounts were imported.") from e
        log.info("bulk student import by %s: created=%d skipped=%d", u.email, created_count, len(errors))
        return {"created": created_count, "skipped": len(errors), "errors": errors}
    finally:
        workbook.close()

@app.get("/api/students")
def list_students(u: User = Depends(staff), s: Session = Depends(db)):
    rows = s.query(User).filter_by(role="student").order_by(User.id.desc()).all()
    return [{"id": r.id, "name": r.name, "email": r.email or "", "department": r.department,
             "created_at": r.created_at.isoformat() + "Z" if r.created_at else ""} for r in rows]

@app.delete("/api/students/{student_id}")
def delete_student(student_id: int, u: User = Depends(staff), s: Session = Depends(db)):
    st = s.get(User, student_id)
    if not st or st.role != "student": raise HTTPException(404, "Student not found")
    # Clean up student's chat history
    s.query(Message).filter_by(student_id=student_id).delete()
    s.query(ChatSession).filter_by(student_id=student_id).delete()
    s.delete(st); s.commit()
    log.info("student deleted by %s: id=%d", u.email, student_id)
    return {"ok": True}

# ── Documents ────────────────────────────────────────────────────────────────
dj = lambda d: {k: getattr(d, k) for k in ("id", "file_name", "file_type", "uploaded_by_name", "uploaded_by_role",
                "department", "status", "error", "page_count", "category", "note")} | {"uploaded_at": d.uploaded_at.isoformat() + "Z"}

def clean_meta(category, note):
    """Validate the staff answers. Returns (category, note) or raises ValueError with a readable reason."""
    category, note = (category or "").strip().lower(), re.sub(r"\s+", " ", note or "").strip()
    if category not in CATEGORIES: raise ValueError("Choose what this document is about")
    if not note: raise ValueError("Add a short note for students")
    if len(note) > NOTE_MAX: raise ValueError(f"Note must be {NOTE_MAX} characters or fewer")
    return category, note

def process(doc_id: int):
    s = Session_(); d = s.get(Document, doc_id)
    try:
        d.status, d.error = "processing", None; s.commit()
        s.query(Chunk).filter_by(document_id=doc_id).delete()
        pages = rag.extract(d.storage_path, d.file_type); d.page_count = len(pages); items = []
        for n, text in enumerate(pages, 1):
            for c in rag.chunk(text):
                ch = Chunk(document_id=doc_id, page_number=n, text=c); s.add(ch); items.append(ch)
        if not items: raise ValueError("No readable text found in this document")
        for ch, v in zip(items, rag.emb([c.text for c in items])): ch.embedding = rag.to_bytes(v)
        d.status = "ready"  # searchable only after successful indexing
        log.info("indexed doc %s (%d chunks)", doc_id, len(items))
    except Exception as e:
        s.rollback(); d = s.get(Document, doc_id); d.status, d.error = "failed", str(e)[:500]
        log.exception("processing failed for doc %s", doc_id)
    finally:
        s.commit(); s.close()

@app.post("/api/documents")
async def upload(bg: BackgroundTasks, files: list[UploadFile] = File(...), meta: str = Form("[]"),
                 u: User = Depends(staff), s: Session = Depends(db)):
    try: answers = json.loads(meta); assert isinstance(answers, list)
    except Exception: raise HTTPException(400, "Invalid upload details")
    out = []
    for i, f in enumerate(files):
        a = answers[i] if i < len(answers) and isinstance(answers[i], dict) else {}
        try: cat, note = clean_meta(a.get("category"), a.get("note"))
        except ValueError as e: out.append({"file": f.filename, "error": str(e)}); continue
        ext = f.filename.rsplit(".", 1)[-1].lower() if "." in f.filename else ""
        if ext not in EXT: out.append({"file": f.filename, "error": "Unsupported file type"}); continue
        data = await f.read()
        if len(data) > MAX_BYTES: out.append({"file": f.filename, "error": "File exceeds 25 MB"}); continue
        path = os.path.join(STORE, f"{uuid.uuid4().hex}.{ext}")
        with open(path, "wb") as fh: fh.write(data)
        d = Document(file_name=os.path.basename(f.filename), file_type=ext, storage_path=path, uploaded_by=u.id,
                     uploaded_by_name=u.name, uploaded_by_role=u.role, department=u.department,
                     category=cat, note=note)
        s.add(d); s.commit(); bg.add_task(process, d.id)
        log.info("upload: %s by %s", d.file_name, u.email); out.append(dj(d))
    return out

@app.get("/api/documents")
def documents(u: User = Depends(staff), s: Session = Depends(db)):
    return [dj(d) for d in s.query(Document).order_by(Document.id.desc())]

class Meta(BaseModel): category: str = ""; note: str = ""

@app.patch("/api/documents/{doc_id}")
def edit_meta(doc_id: int, b: Meta, u: User = Depends(staff), s: Session = Depends(db)):
    d = s.get(Document, doc_id)
    if not d: raise HTTPException(404, "Document not found")
    try: d.category, d.note = clean_meta(b.category, b.note)
    except ValueError as e: raise HTTPException(400, str(e))
    s.commit(); return dj(d)

@app.get("/api/categories")
def categories(u: User = Depends(current)):
    return [{"key": k, "label": v[0], "hint": v[1]} for k, v in CATEGORIES.items()]

@app.get("/api/updates")
def updates(u: User = Depends(current), s: Session = Depends(db)):
    """What is available for students: ready documents with the category and short note staff gave at upload."""
    rows = s.query(Document).filter(Document.status == "ready").order_by(Document.id.desc()).all()
    return [{"id": d.id, "file_name": d.file_name, "category": d.category if d.category in CATEGORIES else "other",
             "note": d.note or "", "uploaded_at": d.uploaded_at.isoformat() + "Z"} for d in rows]

@app.post("/api/documents/{doc_id}/retry")
def retry(doc_id: int, bg: BackgroundTasks, u: User = Depends(staff), s: Session = Depends(db)):
    d = s.get(Document, doc_id)
    if not d or d.status in ("uploaded", "processing"): raise HTTPException(400, "Document is already being processed")
    bg.add_task(process, doc_id); return {"ok": True}

@app.delete("/api/documents/{doc_id}")
def delete_document(doc_id: int, u: User = Depends(staff), s: Session = Depends(db)):
    d = s.get(Document, doc_id)
    if not d:
        raise HTTPException(404, "Document not found")
    if d.status in ("uploaded", "processing"):
        raise HTTPException(409, "Wait for document processing to finish before deleting it")
    path = d.storage_path
    s.query(Chunk).filter_by(document_id=doc_id).delete()
    s.delete(d)
    s.commit()
    try:
        if os.path.exists(path):
            os.remove(path)
    except OSError as e:
        log.exception("document record %s deleted but its file could not be removed", doc_id)
        raise HTTPException(500, "Document was removed from the list, but its stored file could not be deleted") from e
    log.info("document deleted by %s: %s", u.email, doc_id)
    return {"ok": True}

@app.get("/api/documents/{doc_id}/file")
def file(doc_id: int, download: bool = False, u: User = Depends(current), s: Session = Depends(db)):
    d = s.get(Document, doc_id)
    if not d or (u.role != "staff" and d.status != "ready"): raise HTTPException(404, "Document not found")
    log.info("file access: doc %s by %s (download=%s)", doc_id, u.username, download)
    return FileResponse(d.storage_path, filename=d.file_name,
                        content_disposition_type="attachment" if download else "inline")

# ── Chat ─────────────────────────────────────────────────────────────────────
@app.post("/api/chat")
def chat(b: Ask, u: User = Depends(current), s: Session = Depends(db)):
    q = b.question.strip()
    if not q: raise HTTPException(400, "Question is empty")
    cs = s.get(ChatSession, b.session_id) if b.session_id else None
    if b.session_id and (not cs or cs.student_id != u.id): raise HTTPException(404, "Chat not found")
    if not cs:
        cs = ChatSession(student_id=u.id, title=q[:48] + ("…" if len(q) > 48 else "")); s.add(cs); s.flush()
    cs.updated_at = datetime.utcnow()
    s.add(Message(student_id=u.id, session_id=cs.id, role="user", text=q, meta={}))
    rows = (s.query(Chunk.id, Chunk.text, Chunk.page_number, Chunk.embedding, Document.id.label("did"), Document.file_name)
            .join(Document, Document.id == Chunk.document_id).filter(Document.status == "ready").all())
    by = {r.id: r for r in rows}
    corpus = [(r.id, r.text, r.embedding, r.did) for r in rows]
    query = rag.correct_query(q, [r.text for r in rows]) if rows else q
    if query != q:
        log.info("normalized question for retrieval: %r -> %r", q, query)
    hits = rag.search(query, corpus, k=3) if rows else []
    conf = max((h[2] for h in hits), default=0.0)
    log.info("q=%r ready_chunks=%d conf=%.2f thresh=%.2f hits=%s", q, len(rows), conf, rag.THRESH, [(by[c].file_name, by[c].page_number, round(sc, 2)) for c, _, sc in hits])
    answer, meta = NOT_FOUND, {"grounded": False, "confidence": round(conf, 2), "sources": [], "threshold": rag.THRESH,
                    "reason": "low_confidence" if rows else "no_documents"}
    if hits and conf >= rag.THRESH:
        relevant = [hit for hit in hits if hit[2] >= conf - 0.12][:2] or hits[:1]
        ctx = "\n\n".join(
            f"[{by[c].file_name}, page {by[c].page_number}] {rag.excerpt(by[c].text, query)}"
            for c, _, _ in relevant
        )
        try: out = rag.llm(query, ctx)
        except Exception as e:
            log.exception("Ollama generation failed")
            raise HTTPException(503, "The language model is unavailable. Check that Ollama is running with the configured model.") from e
        log.info("LLM replied: %r", out[:150])
        meta["reason"] = "not_in_context"
        if not re.fullmatch(r"NOT[\s_-]*FOUND[.!]?", out.strip(), flags=re.IGNORECASE):
            # Filter out common college/academic boilerplate words so they don't cause false grounding matches
            GENERIC_WORDS = {
                "department", "engineering", "technology", "dindigul", "institution", "academic",
                "autonomous", "office", "student", "students", "faculty", "following", "please",
                "circular", "members", "details", "subject", "semester", "session", "campusflow",
                "assistant", "associate", "professor", "classes", "class", "conducted", "attended",
                "regards", "according", "available", "published", "documents", "course", "college",
                "answer", "teaches", "teach", "hours", "period", "periods", "schedule", "timetable"
            }
            raw_tokens = [re.sub(r"[^a-z0-9]", "", w.lower()) for w in out.split()]
            ans_keywords = {w for w in raw_tokens if len(w) >= 3 and w not in GENERIC_WORDS and w not in rag.STOP}

            # Score each candidate chunk by exact answer keyword matches, then dense score
            scored = []
            seen_keys = set()
            for c, rrf, sc in hits:
                r = by[c]
                key = (r.file_name, r.page_number)
                if key in seen_keys:
                    continue
                seen_keys.add(key)
                chunk_lower = r.text.lower()
                matches = sum(1 for w in ans_keywords if w in chunk_lower)
                scored.append((matches, sc, r))

            scored.sort(key=lambda x: (x[0], x[1]), reverse=True)

            srcs = []
            if scored and scored[0][0] > 0:
                top_m, top_sc, top_r = scored[0]
                srcs.append({"document_id": top_r.did, "file_name": top_r.file_name, "page": top_r.page_number, "score": round(top_sc, 2)})
            elif hits:
                top_hit = hits[0]
                r = by[top_hit[0]]
                srcs.append({"document_id": r.did, "file_name": r.file_name, "page": r.page_number, "score": round(top_hit[2], 2)})

            top = srcs[0]
            share = (f"*CampusFlow Answer*\n{out}\n\nSource: {top['file_name']}\nPage: {top['page']}\n"
                     f"Source / Download: {PUBLIC} (sign in to open)")
            answer, meta = out, {"grounded": True, "confidence": round(conf, 2), "sources": srcs, "share_text": share}
    s.add(Message(student_id=u.id, session_id=cs.id, role="assistant", text=answer, meta=meta)); s.commit()
    return {"role": "assistant", "text": answer, "session_id": cs.id, **meta}

def own_chat(chat_id: int, u: User, s: Session):
    cs = s.get(ChatSession, chat_id)
    if not cs or cs.student_id != u.id: raise HTTPException(404, "Chat not found")
    return cs

@app.get("/api/chats")
def chats(u: User = Depends(current), s: Session = Depends(db)):
    rows = s.query(ChatSession).filter_by(student_id=u.id).order_by(ChatSession.updated_at.desc()).all()
    return [{"id": c.id, "title": c.title, "updated_at": c.updated_at.isoformat() + "Z"} for c in rows]

@app.get("/api/chats/{chat_id}")
def chat_messages(chat_id: int, u: User = Depends(current), s: Session = Depends(db)):
    own_chat(chat_id, u, s)
    ms = s.query(Message).filter_by(session_id=chat_id).order_by(Message.id).all()
    return [{"role": m.role, "text": m.text, **(m.meta or {})} for m in ms]

@app.delete("/api/chats/{chat_id}")
def delete_chat(chat_id: int, u: User = Depends(current), s: Session = Depends(db)):
    cs = own_chat(chat_id, u, s)
    s.query(Message).filter_by(session_id=chat_id).delete(); s.delete(cs); s.commit()
    return {"ok": True}
